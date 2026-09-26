"""EIA-860M "Planned" tab -> generation project records.

EIA-860M (Monthly Update to the Annual Electric Generator Report) lists every generator
that is planned but not yet operating, with plant latitude/longitude and the planned
operation month. It is already structured, so no extraction step is needed.

Generators at one plant that share an in-service month become one project (a 300 MW
solar farm with a 100 MW battery is one site on the map). Each record cites the workbook,
the sheet (a sheet is a "page", as in parsing.py) and the spreadsheet rows it came from.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from app.services.owners import canonical_utility

SOURCE_URL = "https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx"
SHEET = "Planned"
HEADER_ROW = 3  # rows 1-2 are a title and a blank line
# SERTP footprint + Florida, the scope before the load went nationwide.
SOUTHEAST_STATES = ["AL", "GA", "MS", "FL", "TN", "KY", "NC", "SC"]


@dataclass
class PlannedPlant:
    plant_id: int
    plant_name: str
    entity: str
    state: str
    county: str
    balancing_authority: str
    lat: float
    lng: float
    year: int
    month: int
    sheet_index: int
    rows: list[int] = field(default_factory=list)
    generators: list[dict] = field(default_factory=list)

    @property
    def utility(self) -> str:
        return canonical_utility(self.entity) or self.entity

    @property
    def capacity_mw(self) -> float:
        return round(sum(g["mw"] for g in self.generators), 1)

    @property
    def technologies(self) -> list[str]:
        return sorted({g["technology"] for g in self.generators})

    @property
    def name(self) -> str:
        return f"{self.plant_name} ({', '.join(self.technologies)}, {self.capacity_mw:g} MW)"

    @property
    def location_ref(self) -> str:
        county = f"{self.county} County, " if self.county else ""
        return f"{self.plant_name}, {county}{self.state}"

    def excerpt(self) -> str:
        rows = ", ".join(str(r) for r in self.rows)
        lines = [
            f"EIA-860M {SHEET} sheet, row(s) {rows}. Plant {self.plant_id} {self.plant_name} "
            f"({self.entity}), {self.county} County, {self.state}; balancing authority "
            f"{self.balancing_authority}. Planned operation {self.year}-{self.month:02d}."
        ]
        for g in self.generators:
            lines.append(f"Generator {g['id']}: {g['mw']:g} MW {g['technology']}; {g['status']}")
        return "\n".join(lines)[:2000]


def _num(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_planned(path: Path, *, states: list[str] | None = None) -> list[PlannedPlant]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    sheet_index = wb.sheetnames.index(SHEET) + 1
    rows = ws.iter_rows(values_only=True)
    for _ in range(HEADER_ROW - 1):
        next(rows)
    header = [str(h).strip() if h else "" for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}

    plants: dict[tuple[int, int, int], PlannedPlant] = {}
    wanted = {s.upper() for s in states} if states else None
    for row_no, row in enumerate(rows, start=HEADER_ROW + 1):
        cells = dict(zip(col, (row[i] for i in col.values()), strict=True))
        get = cells.get
        plant_id, state = get("Plant ID"), get("Plant State")
        lat, lng = _num(get("Latitude")), _num(get("Longitude"))
        year, month = _num(get("Planned Operation Year")), _num(get("Planned Operation Month"))
        if not isinstance(plant_id, int) or not state or lat is None or lng is None:
            continue  # blank/footer rows
        if wanted and state not in wanted:
            continue
        if year is None or month is None:
            continue
        key = (plant_id, int(year), int(month))
        plant = plants.get(key)
        if plant is None:
            plant = plants[key] = PlannedPlant(
                plant_id=plant_id, plant_name=str(get("Plant Name")).strip(),
                entity=str(get("Entity Name")).strip(), state=state,
                county=str(get("County") or "").strip(),
                balancing_authority=str(get("Balancing Authority Code") or "").strip(),
                lat=lat, lng=lng, year=int(year), month=int(month), sheet_index=sheet_index,
            )
        plant.rows.append(row_no)
        plant.generators.append({
            "id": str(get("Generator ID")),
            "mw": _num(get("Nameplate Capacity (MW)")) or 0.0,
            "technology": str(get("Technology") or "").strip(),
            "status": str(get("Status") or "").strip(),
        })
    wb.close()
    return sorted(plants.values(), key=lambda p: (p.state, p.plant_name, p.year, p.month))


def by_utility(plants: list[PlannedPlant]) -> dict[str, list[PlannedPlant]]:
    grouped: dict[str, list[PlannedPlant]] = defaultdict(list)
    for p in plants:
        grouped[p.utility].append(p)
    return grouped
